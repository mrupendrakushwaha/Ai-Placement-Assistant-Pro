import os
import re
import io
import time
import html
import requests
import streamlit as st
import pdfplumber
from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    PageBreak, KeepTogether
)

# ----------------------------- Page setup -----------------------------
st.set_page_config(
    page_title="AI Placement Assistant Pro",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="expanded",
)

NAVY = "#163A5F"
BLUE = "#2471A3"
TEAL = "#159A9C"
LIGHT_BLUE = "#EAF3F8"
PALE = "#F5F8FB"
DARK = "#23313F"
MUTED = "#667788"
GREEN = "#238636"
AMBER = "#B7791F"
RED = "#C0392B"

st.markdown(
    f"""
    <style>
    .stApp {{ background: #F5F8FB; color: {DARK}; }}
    .block-container {{ padding-top: 1.3rem; padding-bottom: 2rem; max-width: 1200px; }}
    .hero {{
        background: linear-gradient(120deg, {NAVY}, {BLUE});
        padding: 1.35rem 1.5rem; border-radius: 18px; color: white;
        margin-bottom: 1rem; box-shadow: 0 6px 20px rgba(22,58,95,.12);
    }}
    .hero h1 {{ color: white !important; font-size: 2rem !important; margin: 0 0 .35rem 0; font-weight: 800 !important; }}
    .hero p {{ color: #EAF4FB !important; margin: 0; font-size: 1rem; }}
    .section-title {{
        color: {NAVY}; font-weight: 800; font-size: 1.25rem;
        border-left: 5px solid {TEAL}; padding: .15rem 0 .15rem .65rem;
        margin: 1rem 0 .7rem 0;
    }}
    .metric-card {{
        background: white; border: 1px solid #DCE6EE; border-radius: 14px;
        padding: 1rem; min-height: 115px; box-shadow: 0 2px 8px rgba(25,50,75,.04);
    }}
    .metric-label {{ color: {MUTED}; font-size: .88rem; font-weight: 650; }}
    .metric-value {{ color: {NAVY}; font-size: 1.8rem; font-weight: 850; margin-top: .25rem; }}
    .small-note {{ color: {MUTED}; font-size: .85rem; }}
    div[data-testid="stFileUploader"] {{
        background: white; border: 1px dashed #9FB8CA; border-radius: 12px; padding: .6rem;
    }}
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    """
    <div class="hero">
      <h1>🎯 AI Placement Assistant Pro</h1>
      <p>Resume analysis • ATS insights • Skill-gap review • Interview preparation • 3-month action plan</p>
    </div>
    """,
    unsafe_allow_html=True,
)

# ----------------------------- API configuration -----------------------------
def get_secret(name):
    try:
        value = st.secrets.get(name)
        if value:
            return str(value).strip()
    except Exception:
        pass
    return os.environ.get(name, "").strip()

GROQ_API_KEY = get_secret("GROQ_API_KEY")
GROQ_API_URL = "https://api.groq.com/openai/v1/chat/completions"
DEFAULT_GROQ_MODEL = get_secret("GROQ_MODEL") or "openai/gpt-oss-120b"

with st.sidebar:
    st.markdown("## ⚙️ Settings")
    st.markdown("**AI provider:** Groq")
    st.caption(f"Model: `{DEFAULT_GROQ_MODEL}`")
    st.markdown("---")
    st.markdown("### 🔐 API key")
    if GROQ_API_KEY:
        st.success("Groq API key detected")
    else:
        st.warning("Groq API key नहीं मिली। Streamlit Secrets में `GROQ_API_KEY` जोड़ें।")
    st.markdown("---")
    st.caption("Privacy tip: Resume को केवल भरोसेमंद environment में upload करें। API key कभी report या public repository में न डालें।")

# ----------------------------- Resume text extraction -----------------------------
def extract_pdf_text(file_bytes):
    pages = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text.strip():
                pages.append(text)
    return "\n".join(pages).strip()

def extract_docx_text(file_bytes):
    doc = Document(io.BytesIO(file_bytes))
    chunks = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                chunks.append(row_text)
    return "\n".join(chunks).strip()

def extract_resume_text(uploaded):
    raw = uploaded.getvalue()
    suffix = uploaded.name.lower().rsplit(".", 1)[-1]
    if suffix == "pdf":
        return extract_pdf_text(raw)
    if suffix == "docx":
        return extract_docx_text(raw)
    raise ValueError("यह file type supported नहीं है। PDF या DOCX upload करें।")

def normalize_text(text):
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()

def extract_candidate_name(resume_text):
    """Use simple local heuristics first so the name doesn't depend on an AI API call."""
    lines = [re.sub(r"\s+", " ", x).strip(" |•\t") for x in resume_text.splitlines()]
    lines = [x for x in lines if x]
    skip = re.compile(
        r"^(resume|curriculum vitae|cv|profile|summary|objective|contact|"
        r"email|phone|mobile|linkedin|github|portfolio|education|experience|"
        r"skills|projects|certifications|internship|address)\b", re.I
    )
    for line in lines[:12]:
        if skip.search(line):
            continue
        if "@" in line or re.search(r"\+?\d[\d\s().-]{7,}", line):
            continue
        if re.search(r"https?://|www\.|linkedin\.com|github\.com", line, re.I):
            continue
        words = re.findall(r"[A-Za-z][A-Za-z.'’-]*", line)
        if 2 <= len(words) <= 4 and len(line) <= 55:
            if all(len(w) > 1 for w in words) and not any(
                w.lower() in {"engineer", "developer", "student", "analyst", "manager"}
                for w in words
            ):
                return line.title() if line.isupper() else line
    return "Candidate"

# ----------------------------- AI call -----------------------------
def call_groq(prompt, max_tokens=2200, retries=2):
    if not GROQ_API_KEY:
        return None, "Groq API key missing"
    payload = {
        "model": DEFAULT_GROQ_MODEL,
        "messages": [
            {
                "role": "system",
                "content": (
                    "You are a careful career-placement assistant. Analyze only evidence in the supplied resume. "
                    "Do not invent degrees, internships, scores, achievements, or salary guarantees. "
                    "Label uncertain conclusions as estimates. Return clean plain text with clear section headings "
                    "and bullet points. Avoid unusual symbols and emojis."
                ),
            },
            {"role": "user", "content": prompt},
        ],
        "max_tokens": max_tokens,
        "temperature": 0.2,
    }
    headers = {
        "Authorization": f"Bearer {GROQ_API_KEY}",
        "Content-Type": "application/json",
    }
    delay = 1
    for attempt in range(retries + 1):
        try:
            response = requests.post(
                GROQ_API_URL, headers=headers, json=payload, timeout=60
            )
            if response.status_code == 429 and attempt < retries:
                time.sleep(delay)
                delay *= 2
                continue
            if not response.ok:
                try:
                    detail = response.json().get("error", {}).get("message", response.text)
                except Exception:
                    detail = response.text
                return None, f"Groq API HTTP {response.status_code}: {detail[:500]}"
            data = response.json()
            choices = data.get("choices", [])
            if not choices:
                return None, "API response में कोई answer नहीं मिला।"
            content = choices[0].get("message", {}).get("content", "")
            if not content.strip():
                return None, "AI ने खाली report लौटाई।"
            return content.strip(), None
        except requests.RequestException as exc:
            if attempt >= retries:
                return None, f"Network/API request failed: {exc}"
            time.sleep(delay)
            delay *= 2
    return None, "Unknown Groq API error"

def make_analysis_prompt(resume_text, github_username, github_stats):
    return f"""
Analyze the resume for an entry-level job seeker. Be realistic and evidence-based.
Use plain text. Use these exact headings on separate lines:
ATS SCORE
GITHUB SCORE
PLACEMENT READINESS
TOP 5 STRENGTHS
TOP 5 MISSING SKILLS
EXPECTED PACKAGE
INTERVIEW READINESS
TOP 10 HR QUESTIONS
TOP 10 TECHNICAL QUESTIONS
DETAILED 3 MONTH ROADMAP

For each score, give a score out of 10 and 3-5 bullet reasons. Do not imply the ATS score is from a real employer ATS.
Expected Package must be a cautious market estimate, clearly say actual salary varies by role, location, company and skills, and do not promise a salary.
Provide 5 strengths and 5 missing skills, 10 HR questions, 10 technical questions, and at least 5 actionable bullets for each month.
Use only the resume facts. If information is absent, say it is not stated instead of inventing it.
GitHub username: {github_username or "Not provided"}
GitHub public stats (if available): {github_stats}

RESUME TEXT:
{resume_text[:18000]}
"""

# ----------------------------- GitHub stats -----------------------------
def get_github_stats(username):
    if not username:
        return "Not provided"
    if not re.fullmatch(r"[A-Za-z0-9-]{1,39}", username.strip()):
        return "Username format invalid"
    try:
        r = requests.get(
            f"https://api.github.com/users/{username.strip()}",
            headers={"Accept": "application/vnd.github+json"},
            timeout=10,
        )
        if r.status_code == 200:
            user = r.json()
            return (
                f"Public repositories: {user.get('public_repos', 0)}, "
                f"followers: {user.get('followers', 0)}, "
                f"following: {user.get('following', 0)}"
            )
        return f"GitHub stats unavailable (HTTP {r.status_code})"
    except requests.RequestException:
        return "GitHub stats unavailable (network error)"

# ----------------------------- PDF report generation -----------------------------
def clean_pdf_text(text):
    """Normalize model-generated text for ReportLab's standard Helvetica fonts."""
    text = str(text or "")
    replacements = {
        "\u2011": "-", "\u2012": "-", "\u2013": "-", "\u2014": "-", "\u2212": "-",
        "\u2022": "-", "\u00a0": " ", "\u20b9": "INR ",
        "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
        "\u25cf": "-", "\u25aa": "-", "\u2026": "...",
        "\ufffd": "", "\u200b": "", "\u200c": "", "\u200d": "",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)

    # Remove markdown markers that ReportLab would otherwise print literally.
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"__(.*?)__", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\s)(.*?)(?<!\s)\*(?!\*)", r"\1", text)

    # Repair common corrupted separators/spacing seen in AI-generated reports.
    text = re.sub(r"(?i)\bMonth\s*\?\s*([123])\b", r"Month \1", text)
    text = re.sub(r"(?i)\b(\d+)\?\s*(minutes?|hours?|days?|weeks?|months?)\b", r"\1 \2", text)
    text = re.sub(r"(\d+)\?\s*%", r"\1%", text)
    text = re.sub(r"(\d+)\?\s*([kKmM])\b", r"\1\2", text)
    text = re.sub(r"(?i)\bRs\.?\s*(\d+)\?\s*-\s*\??\s*(\d+)\??\s*Lakhs?\b", r"INR \1-\2 Lakhs", text)
    text = re.sub(r"(?i)\b(\d+)\?\s*-\s*\??\s*(\d+)\??\s*Lakhs?\b", r"\1-\2 Lakhs", text)
    text = re.sub(r"(?<=\w)\?(?=\w)", " ", text)
    text = re.sub(r"\s{2,}", " ", text)

    # Standard ReportLab fonts cannot render many Unicode glyphs. Replace only
    # remaining unsupported characters, rather than letting them become boxes.
    text = text.encode("latin-1", "replace").decode("latin-1")
    return text

def safe_paragraph(text, style):
    # Escape AI output before passing it to ReportLab's mini-XML parser.
    return Paragraph(html.escape(clean_pdf_text(text)).replace("\n", "<br/>"), style)

def build_pdf(candidate_name, github_username, report, scores):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, rightMargin=17*mm, leftMargin=17*mm,
        topMargin=18*mm, bottomMargin=18*mm, title="AI Placement Analysis Report",
        author="AI Placement Assistant Pro",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="HeroTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=21, leading=25, textColor=colors.white, alignment=TA_CENTER,
        spaceAfter=5,
    ))
    styles.add(ParagraphStyle(
        name="HeroSub", parent=styles["Normal"], fontName="Helvetica",
        fontSize=9, leading=13, textColor=colors.HexColor("#EAF4FB"), alignment=TA_CENTER,
    ))
    styles.add(ParagraphStyle(
        name="SectionHeading", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=12.5, leading=16, textColor=colors.HexColor(NAVY),
        spaceBefore=10, spaceAfter=6, keepWithNext=True,
    ))
    styles.add(ParagraphStyle(
        name="SubHeadingCustom", parent=styles["Heading3"], fontName="Helvetica-Bold",
        fontSize=10, leading=13, textColor=colors.HexColor(BLUE),
        spaceBefore=5, spaceAfter=3, keepWithNext=True,
    ))
    styles.add(ParagraphStyle(
        name="BodyCustom", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=8.7, leading=12, textColor=colors.HexColor(DARK), spaceAfter=3,
    ))
    styles.add(ParagraphStyle(
        name="SmallCustom", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=8, leading=10, textColor=colors.HexColor(MUTED),
    ))

    story = []
    hero = Table(
        [[Paragraph("AI PLACEMENT ASSISTANT PRO", styles["HeroTitle"])],
         [Paragraph("PROFESSIONAL PLACEMENT ANALYSIS REPORT", styles["HeroSub"])]],
        colWidths=[176*mm],
    )
    hero.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor(NAVY)),
        ("BOX", (0,0), (-1,-1), 0, colors.HexColor(NAVY)),
        ("LEFTPADDING", (0,0), (-1,-1), 12),
        ("RIGHTPADDING", (0,0), (-1,-1), 12),
        ("TOPPADDING", (0,0), (-1,0), 15),
        ("BOTTOMPADDING", (0,0), (-1,0), 2),
        ("TOPPADDING", (0,1), (-1,1), 2),
        ("BOTTOMPADDING", (0,1), (-1,1), 15),
    ]))
    story += [hero, Spacer(1, 10)]

    profile_data = [
        [Paragraph("<b>CANDIDATE</b>", styles["SmallCustom"]),
         Paragraph("<b>GITHUB USERNAME</b>", styles["SmallCustom"])],
        [safe_paragraph(candidate_name, styles["BodyCustom"]),
         safe_paragraph(github_username or "Not provided", styles["BodyCustom"])],
    ]
    profile = Table(profile_data, colWidths=[88*mm, 88*mm])
    profile.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), colors.HexColor(PALE)),
        ("BOX", (0,0), (-1,-1), .5, colors.HexColor("#DCE6EE")),
        ("INNERGRID", (0,0), (-1,-1), .35, colors.HexColor("#DCE6EE")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 9), ("RIGHTPADDING", (0,0), (-1,-1), 9),
        ("TOPPADDING", (0,0), (-1,-1), 7), ("BOTTOMPADDING", (0,0), (-1,-1), 7),
    ]))
    story += [profile, Spacer(1, 12)]

    if scores:
        cards = []
        for label, value in scores:
            cards.append(Paragraph(
                f'<font color="{NAVY}"><b>{html.escape(label)}</b></font><br/>'
                f'<font size="18" color="{TEAL}"><b>{value}/10</b></font>',
                styles["BodyCustom"],
            ))
        score_table = Table([cards], colWidths=[176*mm/max(1, len(cards))]*len(cards))
        score_table.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,-1), colors.HexColor(LIGHT_BLUE)),
            ("BOX", (0,0), (-1,-1), .5, colors.HexColor("#D1E2EC")),
            ("INNERGRID", (0,0), (-1,-1), .5, colors.white),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ("ALIGN", (0,0), (-1,-1), "CENTER"),
            ("LEFTPADDING", (0,0), (-1,-1), 5), ("RIGHTPADDING", (0,0), (-1,-1), 5),
            ("TOPPADDING", (0,0), (-1,-1), 9), ("BOTTOMPADDING", (0,0), (-1,-1), 9),
        ]))
        story += [score_table, Spacer(1, 10)]

    known_headings = {
        "ATS SCORE", "GITHUB SCORE", "PLACEMENT READINESS", "TOP 5 STRENGTHS",
        "TOP 5 MISSING SKILLS", "EXPECTED PACKAGE", "INTERVIEW READINESS",
        "TOP 10 HR QUESTIONS", "TOP 10 TECHNICAL QUESTIONS", "DETAILED 3 MONTH ROADMAP",
        "MONTH 1", "MONTH 2", "MONTH 3", "REASONING", "SCORE",
    }
    for raw_line in report.splitlines():
        line = raw_line.strip().lstrip("#").strip()
        if not line:
            story.append(Spacer(1, 3))
            continue
        line = clean_pdf_text(line)
        normalized = re.sub(r"[*_]+", "", line).strip().rstrip(":").upper()
        if normalized in known_headings or (
            len(line) < 48 and line.isupper() and any(c.isalpha() for c in line)
        ):
            story.append(Paragraph(html.escape(line.rstrip(":")), styles["SectionHeading"]))
        elif re.match(r"^(month\s+[123](?:\s*[-:–].*)?|reasoning|score)\s*:?\s*$", line, re.I):
            story.append(Paragraph(html.escape(line), styles["SubHeadingCustom"]))
        elif line.startswith(("-", "•", "*")):
            story.append(Paragraph("&#8226; " + html.escape(clean_pdf_text(line.lstrip("-•* ").strip())), styles["BodyCustom"]))
        else:
            story.append(safe_paragraph(line, styles["BodyCustom"]))

    story += [Spacer(1, 14)]
    footer = Table(
        [[Paragraph("<b>AI Placement Assistant Pro</b>", styles["SmallCustom"]),
          Paragraph("Career guidance is an estimate, not a job or salary guarantee.", styles["SmallCustom"])]],
        colWidths=[70*mm, 106*mm],
    )
    footer.setStyle(TableStyle([
        ("LINEABOVE", (0,0), (-1,-1), .8, colors.HexColor(TEAL)),
        ("TOPPADDING", (0,0), (-1,-1), 7),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
    ]))
    story.append(footer)

    def add_page(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#DCE6EE"))
        canvas.line(17*mm, 13*mm, A4[0]-17*mm, 13*mm)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor(MUTED))
        canvas.drawString(17*mm, 8*mm, "AI Placement Assistant Pro")
        canvas.drawRightString(A4[0]-17*mm, 8*mm, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=add_page, onLaterPages=add_page)
    buffer.seek(0)
    return buffer.getvalue()

def parse_scores(report):
    scores = []
    for label in ["ATS SCORE", "GITHUB SCORE", "PLACEMENT READINESS", "INTERVIEW READINESS"]:
        pattern = rf"{re.escape(label)}[\s\S]{{0,180}}?(?:score\s*:?\s*)?(\d{{1,2}})\s*/\s*10"
        match = re.search(pattern, report, re.I)
        if match:
            value = min(10, int(match.group(1)))
            short = {
                "ATS SCORE": "ATS",
                "GITHUB SCORE": "GitHub",
                "PLACEMENT READINESS": "Placement",
                "INTERVIEW READINESS": "Interview",
            }[label]
            scores.append((short, value))
    return scores

# ----------------------------- Input UI -----------------------------
st.markdown('<div class="section-title">1. Upload your resume</div>', unsafe_allow_html=True)
left, right = st.columns([1.3, 1])
with left:
    uploaded_file = st.file_uploader(
        "Resume file (PDF or DOCX)",
        type=["pdf", "docx"],
        help="PDF and DOCX supported. Maximum file size: 20 MB.",
    )
    st.caption("Supported: text-based PDF and DOCX. Scanned/image-only PDFs may need OCR.")
with right:
    github_username = st.text_input(
        "GitHub username (optional)",
        placeholder="e.g. mrupendrakushwaha",
    )
    st.caption("Public GitHub profile stats are optional and do not require a token.")

if uploaded_file and uploaded_file.size > 20 * 1024 * 1024:
    st.error("File 20 MB से बड़ी है। कृपया छोटी PDF/DOCX upload करें।")
    st.stop()

if uploaded_file:
    st.success(f"Selected: {uploaded_file.name} · {uploaded_file.size / 1024:.0f} KB")

analyze_clicked = st.button("✨ Analyze Resume", type="primary", use_container_width=True, disabled=not uploaded_file)

if analyze_clicked:
    try:
        with st.spinner("Resume text पढ़ रहे हैं..."):
            resume_text = normalize_text(extract_resume_text(uploaded_file))
    except Exception as exc:
        st.error(f"Resume पढ़ने में समस्या हुई: {exc}")
        st.stop()

    if not resume_text:
        st.error(
            "इस फाइल से text नहीं मिला। अगर यह scanned/image-only PDF है, तो OCR से text निकालकर DOCX या searchable PDF upload करें।"
        )
        st.stop()

    candidate_name = extract_candidate_name(resume_text)
    github_stats = get_github_stats(github_username.strip())
    prompt = make_analysis_prompt(resume_text, github_username.strip(), github_stats)

    with st.spinner("AI आपकी placement report बना रहा है..."):
        report, error = call_groq(prompt)

    if error:
        st.error(f"AI report नहीं बन पाई: {error}")
        st.info("जाँचें कि Streamlit Secrets में GROQ_API_KEY सही है और चुना गया model Groq Console में उपलब्ध है।")
        st.stop()

    # Save result in session state so download remains available after reruns.
    st.session_state["placement_result"] = {
        "candidate_name": candidate_name,
        "github_username": github_username.strip(),
        "report": report,
        "scores": parse_scores(report),
        "filename": uploaded_file.name,
    }

# ----------------------------- Results dashboard -----------------------------
result = st.session_state.get("placement_result")
if result:
    st.markdown('<div class="section-title">2. Placement dashboard</div>', unsafe_allow_html=True)
    st.markdown(
        f"**Candidate:** {result['candidate_name']}  ·  **Resume:** {result['filename']}"
    )

    scores = result["scores"]
    if scores:
        cols = st.columns(len(scores))
        for col, (label, score) in zip(cols, scores):
            with col:
                st.markdown(
                    f'<div class="metric-card"><div class="metric-label">{html.escape(label)} SCORE</div>'
                    f'<div class="metric-value">{score}<span style="font-size:1rem;color:{MUTED}"> / 10</span></div></div>',
                    unsafe_allow_html=True,
                )
                st.progress(score / 10)
        st.markdown('<div class="section-title">Score comparison</div>', unsafe_allow_html=True)
        chart_data = {"Category": [x[0] for x in scores], "Score out of 10": [x[1] for x in scores]}
        st.bar_chart(chart_data, x="Category", y="Score out of 10")
    else:
        st.info("Report में score format नहीं मिला; detailed analysis नीचे उपलब्ध है।")

    st.markdown('<div class="section-title">3. Detailed placement report</div>', unsafe_allow_html=True)
    # Render each line safely; heading lines get bold colored styling.
    heading_names = {
        "ATS SCORE", "GITHUB SCORE", "PLACEMENT READINESS", "TOP 5 STRENGTHS",
        "TOP 5 MISSING SKILLS", "EXPECTED PACKAGE", "INTERVIEW READINESS",
        "TOP 10 HR QUESTIONS", "TOP 10 TECHNICAL QUESTIONS", "DETAILED 3 MONTH ROADMAP",
        "MONTH 1", "MONTH 2", "MONTH 3",
    }
    for line in result["report"].splitlines():
        clean = line.strip()
        if not clean:
            st.write("")
            continue
        normalized = re.sub(r"[*_]+", "", clean).strip().rstrip(":").upper()
        if normalized in heading_names:
            st.markdown(
                f'<div style="color:{NAVY};font-weight:800;font-size:1.12rem;'
                f'border-left:5px solid {TEAL};padding:.2rem 0 .2rem .6rem;margin:.8rem 0 .35rem 0">'
                f'{html.escape(clean.rstrip(":"))}</div>',
                unsafe_allow_html=True,
            )
        elif clean.startswith(("-", "•", "*")):
            st.markdown("- " + clean.lstrip("-•* ").strip())
        else:
            st.write(clean)

    try:
        pdf_bytes = build_pdf(
            result["candidate_name"], result["github_username"],
            result["report"], result["scores"]
        )
        st.download_button(
            "📥 Download professional PDF report",
            data=pdf_bytes,
            file_name="AI_Placement_Report.pdf",
            mime="application/pdf",
            use_container_width=True,
        )
    except Exception as exc:
        st.error(f"PDF बनाते समय समस्या हुई: {exc}")

st.markdown("---")
st.caption("AI-generated scores are guidance only. Verify important career, salary, and hiring decisions independently.")
